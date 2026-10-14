#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Suite de testes de corrigir_planilha.py.

Valida, item por item, que cada inconsistencia apontada no diagnostico foi de
fato corrigida na planilha gerada - e, igualmente importante, que o script nao
corrigiu demais: duplicata suspeita nao pode sumir calada, receita ausente nao
pode ser inventada e o arquivo de origem nao pode ser tocado.

A planilha de teste e gerada do zero num diretorio temporario a cada execucao, a
partir da planilha original. Os testes nunca leem o arquivo ja existente em
data/, para que um resultado antigo no disco jamais faca um teste passar.

Usa apenas unittest da biblioteca padrao. Nao requer pytest.

Uso
---
    python3 testar_correcoes.py           # saida detalhada, um teste por linha
    python3 testar_correcoes.py -q        # so o resumo
    python3 testar_correcoes.py TestDiagnostico
    python3 testar_correcoes.py TestDiagnostico.test_06_duplicata_exata_removida
"""

from __future__ import annotations

import hashlib
import importlib.util
import shutil
import subprocess
import sys
import tempfile
import unittest
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent
SCRIPT = RAIZ / "corrigir_planilha.py"
ORIGINAL = RAIZ / "data" / "planilha_vendas.xlsx"


def _carregar_modulo():
    if not SCRIPT.exists():
        sys.exit(f"Script nao encontrado: {SCRIPT}")
    spec = importlib.util.spec_from_file_location("corrigir_planilha", SCRIPT)
    modulo = importlib.util.module_from_spec(spec)
    # Registrar antes de exec_module: no Python 3.9 o @dataclass resolve
    # cls.__module__ via sys.modules e quebra sem isso.
    sys.modules["corrigir_planilha"] = modulo
    spec.loader.exec_module(modulo)
    return modulo


cp = _carregar_modulo()


# ===========================================================================
# VALORES ANCORA
# Do diagnostico da planilha original. Mudou um? Ou a origem mudou, ou houve regressao.
# ===========================================================================

LINHAS_ORIGINAIS = 30
LINHAS_ESPERADAS = 29
RECEITA_GANHA = 113_720.00
RECEITA_PERDIDA = 10_090.00      # inclui o pedido cujo status estava em Observacoes
TOTAL_BRUTO = 128_130.00         # soma de todas as linhas com receita, na origem

REGIOES = {"Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"}
STATUS = {"Fechado Ganho", "Fechado Perdido"}

PARES_SUSPEITOS = [(1004, 1017), (1006, 1029), (1007, 1025), (1008, 1014)]
PAR_EXATO = (1001, 1005)         # mantem o menor ID, remove o outro
SEM_RECEITA = {1009, 1021}       # nao podem ser preenchidos por adivinhacao


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


class Planilha:
    """Leitura tipada da planilha gerada, para os testes falarem de negocio."""

    COLUNAS = ["id", "cliente", "regiao", "representante", "data", "receita",
               "produto", "status", "observacoes", "situacao", "pendencias"]

    def __init__(self, caminho: Path) -> None:
        self.caminho = caminho
        self.livro = openpyxl.load_workbook(caminho)
        self.abas = self.livro.sheetnames
        self.ws = self.livro["pedidos_vendas"]
        self.cabecalho = [c.value for c in self.ws[1]]
        self.pedidos = [
            dict(zip(self.COLUNAS, linha))
            for linha in self.ws.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in linha)
        ]
        self.revisar = self._linhas("revisar", ["id", "tipo", "detalhe", "acao"])
        self.log = self._linhas("log_correcoes", ["id", "coluna", "antes", "depois", "regra"])
        self.removidos = self._linhas(
            "removidos",
            ["id", "cliente", "regiao", "representante", "data", "receita",
             "produto", "status", "motivo"])

    def _linhas(self, aba: str, colunas: list[str]) -> list[dict]:
        if aba not in self.livro.sheetnames:
            return []
        ws = self.livro[aba]
        return [dict(zip(colunas, linha))
                for linha in ws.iter_rows(min_row=2, values_only=True)
                if any(v is not None for v in linha)]

    def ids(self) -> set:
        return {p["id"] for p in self.pedidos}

    def por_id(self, pedido_id):
        for p in self.pedidos:
            if p["id"] == pedido_id:
                return p
        return None

    def coluna(self, nome: str) -> list:
        return [p[nome] for p in self.pedidos]

    def pendencias_de(self, pedido_id) -> list[str]:
        return [r["tipo"] for r in self.revisar if r["id"] == pedido_id]

    def formato(self, linha: int, coluna: int) -> str:
        return self.ws.cell(linha, coluna).number_format


class BaseComPlanilha(unittest.TestCase):
    """Gera a planilha uma vez e compartilha entre os testes da classe."""

    gerada: Planilha
    tmp: Path
    hash_original: str

    @classmethod
    def setUpClass(cls) -> None:
        if not ORIGINAL.exists():
            raise unittest.SkipTest(f"Planilha original ausente: {ORIGINAL}")
        cls.hash_original = sha256(ORIGINAL)
        cls._dir = tempfile.mkdtemp(prefix="teste_correcao_")
        cls.tmp = Path(cls._dir)
        destino = cls.tmp / "gerada.xlsx"

        pedidos = cp.carregar(ORIGINAL, cp.ABA_ORIGEM)
        cls.lidos = len(pedidos)
        pedidos, aud = cp.processar(pedidos)
        cp.gravar(destino, pedidos, aud, ORIGINAL, cls.lidos)

        cls.pedidos_memoria = pedidos
        cls.auditoria = aud
        cls.gerada = Planilha(destino)

    @classmethod
    def tearDownClass(cls) -> None:
        shutil.rmtree(cls._dir, ignore_errors=True)


# ===========================================================================
# 1. REGRAS DE NORMALIZACAO (funcoes puras, sem arquivo)
# ===========================================================================

class TestRegrasDeNormalizacao(unittest.TestCase):

    def test_receita_formato_en_us(self):
        """ler_receita: 'R$ 4320.00' (ponto decimal) vira 4320.00"""
        self.assertEqual(cp.ler_receita("R$ 4320.00"), 4320.00)
        self.assertEqual(cp.ler_receita("R$ 1,150.00"), 1150.00)
        self.assertEqual(cp.ler_receita("R$ 990"), 990.00)

    def test_receita_formato_pt_br(self):
        """ler_receita: tambem aceita 'R$ 4.320,00' (padrao brasileiro)"""
        self.assertEqual(cp.ler_receita("R$ 4.320,00"), 4320.00)
        self.assertEqual(cp.ler_receita("R$ 12.400,50"), 12400.50)

    def test_receita_valores_invalidos(self):
        """ler_receita: vazio, nulo e texto sem digito viram None, nao zero"""
        for entrada in (None, "", "   ", "a combinar", "R$"):
            with self.subTest(entrada=entrada):
                self.assertIsNone(cp.ler_receita(entrada))

    def test_receita_numero_passa_direto(self):
        """ler_receita: um numero ja gravado como numero e preservado"""
        self.assertEqual(cp.ler_receita(4320), 4320.00)
        self.assertEqual(cp.ler_receita(4320.5), 4320.50)

    def test_capitalizar_caixa(self):
        """capitalizar: MAIUSCULA e minuscula convergem para a mesma forma"""
        self.assertEqual(cp.capitalizar("maria souza"), "Maria Souza")
        self.assertEqual(cp.capitalizar("PEDRO OLIVEIRA"), "Pedro Oliveira")
        self.assertEqual(cp.capitalizar("DÉRICO MENDES"), "Dérico Mendes")

    def test_capitalizar_hifen(self):
        """capitalizar: hifen tem a parte seguinte capitalizada"""
        self.assertEqual(cp.capitalizar("centro-oeste"), "Centro-Oeste")
        self.assertEqual(cp.capitalizar("CENTRO-OESTE"), "Centro-Oeste")

    def test_capitalizar_particulas(self):
        """capitalizar: particulas ficam minusculas no meio do nome"""
        self.assertEqual(cp.capitalizar("joão da silva"), "João da Silva")
        self.assertEqual(cp.capitalizar("MARIA DOS SANTOS"), "Maria dos Santos")

    def test_capitalizar_nulo(self):
        """capitalizar: nulo e string vazia continuam nulos"""
        self.assertIsNone(cp.capitalizar(None))
        self.assertIsNone(cp.capitalizar("   "))

    def test_chave_ignora_acento_e_caixa(self):
        """chave: acento e caixa nao distinguem dois valores iguais"""
        self.assertEqual(cp.chave("Tomás Bastos"), cp.chave("TOMAS BASTOS"))
        self.assertEqual(cp.chave("Emília  Dias"), cp.chave("emilia dias"))

    def test_chave_regiao_ignora_hifen(self):
        """chave_regiao: 'Centro-Oeste', 'centro oeste' e 'CENTROOESTE' colidem"""
        formas = ["Centro-Oeste", "centro oeste", "CENTROOESTE", "centro-oeste"]
        self.assertEqual(len({cp.chave_regiao(f) for f in formas}), 1)

    def test_ler_data_tipos(self):
        """ler_data: datetime, date e texto em varios formatos viram date"""
        esperado = date(2024, 1, 15)
        self.assertEqual(cp.ler_data(datetime(2024, 1, 15, 0, 0)), esperado)
        self.assertEqual(cp.ler_data(date(2024, 1, 15)), esperado)
        self.assertEqual(cp.ler_data("2024-01-15"), esperado)
        self.assertEqual(cp.ler_data("15/01/2024"), esperado)
        self.assertIsNone(cp.ler_data("data ilegivel"))

    def test_dominios_sao_fechados(self):
        """Os dominios de regiao e status cobrem exatamente os valores previstos"""
        self.assertEqual(set(cp.REGIOES_CANONICAS.values()), REGIOES)
        self.assertEqual(set(cp.STATUS_CANONICO.values()), STATUS)


# ===========================================================================
# 2. UM TESTE POR ITEM DO DIAGNOSTICO
# ===========================================================================

class TestDiagnostico(BaseComPlanilha):

    def test_01_receita_e_numero(self):
        """Item 1: Receita deixou de ser texto e virou numero"""
        valores = [p["receita"] for p in self.gerada.pedidos]
        preenchidos = [v for v in valores if v is not None]
        self.assertTrue(preenchidos, "nenhuma receita foi preenchida")
        for v in preenchidos:
            self.assertIsInstance(v, (int, float), f"receita ainda nao numerica: {v!r}")
            self.assertNotIsInstance(v, str)

    def test_02_receita_tem_formato_de_moeda(self):
        """Item 2: o formato monetario vive na celula, nao no conteudo"""
        coluna = self.gerada.cabecalho.index("Receita") + 1
        for linha in range(2, len(self.gerada.pedidos) + 2):
            self.assertEqual(self.gerada.formato(linha, coluna), cp.FORMATO_MOEDA)

    def test_03_caixa_padronizada(self):
        """Item 3: nenhum valor aparece em duas grafias que so diferem na caixa"""
        for campo in ("cliente", "regiao", "representante", "status"):
            grupos = defaultdict(set)
            for valor in self.gerada.coluna(campo):
                if valor:
                    grupos[cp.chave(valor)].add(valor)
            duplicadas = {k: v for k, v in grupos.items() if len(v) > 1}
            self.assertEqual(duplicadas, {}, f"{campo} ainda tem grafias concorrentes")

    def test_04_regiao_no_dominio(self):
        """Item 4: toda regiao pertence ao dominio fechado de cinco valores"""
        presentes = {v for v in self.gerada.coluna("regiao") if v}
        self.assertTrue(presentes <= REGIOES, f"fora do dominio: {presentes - REGIOES}")

    def test_05_representante_canonico(self):
        """Item 5: cada representante aparece sob uma unica grafia"""
        nomes = [v for v in self.gerada.coluna("representante") if v]
        por_chave = defaultdict(set)
        for nome in nomes:
            por_chave[cp.chave(nome)].add(nome)
        for chave, formas in por_chave.items():
            self.assertEqual(len(formas), 1, f"'{chave}' aparece como {formas}")

    def test_06_duplicata_exata_removida(self):
        """Item 6: a duplicata exata saiu da aba principal e foi para removidos"""
        mantido, removido = PAR_EXATO
        self.assertIn(mantido, self.gerada.ids(), "o original foi removido por engano")
        self.assertNotIn(removido, self.gerada.ids(), "a copia continua na aba principal")

        quarentena = {r["id"] for r in self.gerada.removidos}
        self.assertIn(removido, quarentena, "a copia nao foi preservada em removidos")

        linha = next(r for r in self.gerada.removidos if r["id"] == removido)
        self.assertIn(str(mantido), str(linha["motivo"]),
                      "o motivo nao aponta qual pedido foi mantido")

    def test_07_status_recuperado_de_observacoes(self):
        """Item 7: o status que estava em Observacoes voltou para a coluna Status"""
        for p in self.gerada.pedidos:
            self.assertIsNotNone(p["status"], f"pedido {p['id']} sem status")
            obs = cp.chave(p["observacoes"])
            self.assertNotIn(obs, cp.STATUS_CANONICO,
                             f"pedido {p['id']} ainda tem status em Observacoes")

    def test_08_status_traduzido(self):
        """Item 8: nenhum status em ingles sobrou"""
        presentes = {v for v in self.gerada.coluna("status") if v}
        self.assertTrue(presentes <= STATUS, f"fora do dominio: {presentes - STATUS}")
        for valor in presentes:
            self.assertNotIn("closed", valor.lower())

    def test_09_regiao_imputada(self):
        """Item 9: nao sobrou regiao vazia, e a imputada bate com o territorio"""
        vazias = [p["id"] for p in self.gerada.pedidos if not p["regiao"]]
        self.assertEqual(vazias, [], f"regiao vazia nos pedidos {vazias}")

        territorio = defaultdict(set)
        for p in self.gerada.pedidos:
            territorio[p["representante"]].add(p["regiao"])
        for linha in self.gerada.log:
            if linha["coluna"] == "Regiao" and linha["antes"] == "(vazio)":
                pedido = self.gerada.por_id(linha["id"])
                self.assertIn(pedido["regiao"], territorio[pedido["representante"]])

    def test_10_representante_imputado(self):
        """Item 10: nao sobrou representante vazio"""
        vazios = [p["id"] for p in self.gerada.pedidos if not p["representante"]]
        self.assertEqual(vazios, [], f"representante vazio nos pedidos {vazios}")

    def test_11_id_e_inteiro(self):
        """Item 11: o ID do pedido deixou de ser decimal"""
        # Checar so o arquivo nao basta: ao reler 1001.0 o openpyxl devolve int e
        # mascara a falta da conversao. O objeto em memoria e que e autoritativo.
        for pedido in self.pedidos_memoria:
            self.assertIsInstance(pedido.id, int,
                                  f"ID ainda decimal em memoria: {pedido.id!r}")
            self.assertNotIsInstance(pedido.id, float)
        for p in self.gerada.pedidos:
            self.assertIsInstance(p["id"], int, f"ID ainda decimal: {p['id']!r}")
        coluna = self.gerada.cabecalho.index("ID Pedido") + 1
        self.assertEqual(self.gerada.formato(2, coluna), "0")

    def test_11b_conversao_de_id_registrada_no_log(self):
        """Item 11: toda conversao de ID decimal aparece no log de correcoes

        Guarda uma regressao ja vista: como 1001.0 == 1001 e verdadeiro em
        Python, uma comparacao sem tipo faria a conversao sumir do log.
        """
        origem = openpyxl.load_workbook(ORIGINAL)[cp.ABA_ORIGEM]
        decimais = {int(l[0]) for l in origem.iter_rows(min_row=2, values_only=True)
                    if isinstance(l[0], float)}
        if not decimais:
            self.skipTest("a origem nao tem ID decimal")
        registrados = {l["id"] for l in self.gerada.log if l["coluna"] == "ID Pedido"}
        self.assertEqual(decimais, registrados,
                         f"conversoes de ID fora do log: {decimais - registrados}")

    def test_12_observacoes_sem_marcacao_de_auditoria(self):
        """Item 12: Observacoes guarda nota de negocio, nao marcacao de QA"""
        for p in self.gerada.pedidos:
            if not p["observacoes"]:
                continue
            for trecho in str(p["observacoes"]).split(";"):
                k = cp.chave(trecho)
                self.assertNotIn(k, cp.MARCACOES_DE_AUDITORIA,
                                 f"pedido {p['id']} mantem a marcacao '{trecho.strip()}'")
                self.assertNotIn(k, cp.STATUS_CANONICO)

    def test_13_duplicata_suspeita_sinalizada_e_preservada(self):
        """Item 13: suspeita e sinalizada para revisao, nunca removida em silencio"""
        for a, b in PARES_SUSPEITOS:
            with self.subTest(par=(a, b)):
                self.assertIn(a, self.gerada.ids(), f"pedido {a} sumiu sem decisao humana")
                self.assertIn(b, self.gerada.ids(), f"pedido {b} sumiu sem decisao humana")
                for pedido_id, parceiro in ((a, b), (b, a)):
                    self.assertIn("Duplicata suspeita", self.gerada.pendencias_de(pedido_id),
                                  f"pedido {pedido_id} nao foi sinalizado")
                    detalhe = next(r["detalhe"] for r in self.gerada.revisar
                                   if r["id"] == pedido_id and r["tipo"] == "Duplicata suspeita")
                    self.assertIn(str(parceiro), str(detalhe),
                                  "o detalhe nao aponta o pedido par")

    def test_14_receita_ausente_nao_foi_inventada(self):
        """Item 14: receita que nao existe continua vazia e vira pendencia"""
        for pedido_id in SEM_RECEITA:
            with self.subTest(pedido=pedido_id):
                pedido = self.gerada.por_id(pedido_id)
                self.assertIsNotNone(pedido, f"pedido {pedido_id} desapareceu")
                self.assertIsNone(pedido["receita"],
                                  "um valor foi inventado para uma receita ausente")
                self.assertIn("Receita ausente", self.gerada.pendencias_de(pedido_id))

    def test_15_id_fora_de_ordem_sinalizado(self):
        """Item 15: ID cuja data retrocede e apontado para conferencia"""
        tipos = {r["tipo"] for r in self.gerada.revisar}
        self.assertIn("ID fora de ordem cronologica", tipos)

        numerados = sorted((p for p in self.gerada.pedidos if p["data"]), key=lambda p: p["id"])
        quebras = {b["id"] for a, b in zip(numerados, numerados[1:]) if b["data"] < a["data"]}
        sinalizados = {r["id"] for r in self.gerada.revisar
                       if r["tipo"] == "ID fora de ordem cronologica"}
        self.assertEqual(quebras, sinalizados, "ha quebra de cronologia nao sinalizada")

    def test_16_representante_em_duas_regioes_alertado(self):
        """Item 16: acumulo de territorio vira alerta, nao correcao automatica"""
        territorio = defaultdict(set)
        for p in self.gerada.pedidos:
            if p["representante"] and p["regiao"]:
                territorio[p["representante"]].add(p["regiao"])
        acumulam = {rep for rep, regioes in territorio.items() if len(regioes) > 1}

        self.assertEqual(len(self.auditoria.alertas), len(acumulam),
                         "a quantidade de alertas nao bate com os representantes em 2 regioes")
        for rep in acumulam:
            self.assertTrue(any(rep in alerta for alerta in self.auditoria.alertas),
                            f"{rep} atua em varias regioes e nao gerou alerta")


# ===========================================================================
# 3. INTEGRIDADE - o script nao pode perder nem inventar dado
# ===========================================================================

class TestIntegridade(BaseComPlanilha):

    def test_arquivo_original_intacto(self):
        """O arquivo de origem nao foi modificado, byte a byte"""
        self.assertEqual(sha256(ORIGINAL), self.hash_original)

    def test_nenhuma_linha_perdida(self):
        """Toda linha da origem esta na aba principal ou na aba removidos"""
        origem = openpyxl.load_workbook(ORIGINAL)[cp.ABA_ORIGEM]
        ids_origem = {int(linha[0]) for linha in origem.iter_rows(min_row=2, values_only=True)
                      if linha[0] is not None}
        ids_saida = self.gerada.ids() | {r["id"] for r in self.gerada.removidos}
        self.assertEqual(ids_origem, ids_saida,
                         f"diferenca de IDs: {ids_origem ^ ids_saida}")

    def test_contagem_de_linhas(self):
        """A contagem bate: lidas = mantidas + removidas"""
        self.assertEqual(self.lidos, LINHAS_ORIGINAIS)
        self.assertEqual(len(self.gerada.pedidos), LINHAS_ESPERADAS)
        self.assertEqual(len(self.gerada.pedidos) + len(self.gerada.removidos),
                         LINHAS_ORIGINAIS)

    def test_receita_conservada(self):
        """Nenhum centavo apareceu ou sumiu: mantidos + removidos = total da origem"""
        origem = openpyxl.load_workbook(ORIGINAL)[cp.ABA_ORIGEM]
        bruto = sum(cp.ler_receita(linha[5]) or 0
                    for linha in origem.iter_rows(min_row=2, values_only=True))
        self.assertAlmostEqual(bruto, TOTAL_BRUTO, places=2)

        mantido = sum(p["receita"] or 0 for p in self.gerada.pedidos)
        removido = sum(r["receita"] or 0 for r in self.gerada.removidos)
        self.assertAlmostEqual(mantido + removido, bruto, places=2)

    def test_receita_ganha_bate_com_diagnostico(self):
        """A receita ganha e a perdida reproduzem os numeros do diagnostico"""
        ganha = sum(p["receita"] or 0 for p in self.gerada.pedidos
                    if p["status"] == "Fechado Ganho")
        perdida = sum(p["receita"] or 0 for p in self.gerada.pedidos
                      if p["status"] == "Fechado Perdido")
        self.assertAlmostEqual(ganha, RECEITA_GANHA, places=2)
        self.assertAlmostEqual(perdida, RECEITA_PERDIDA, places=2)

    def test_log_cobre_toda_alteracao(self):
        """Cada diferenca entre origem e saida tem uma linha no log de correcoes"""
        origem = {int(l[0]): l for l in
                  openpyxl.load_workbook(ORIGINAL)[cp.ABA_ORIGEM].iter_rows(min_row=2, values_only=True)
                  if l[0] is not None}
        mapa = {"Nome do Cliente": 1, "Regiao": 2, "Representante": 3, "Status": 7}
        registradas = {(l["id"], l["coluna"]) for l in self.gerada.log}

        for p in self.gerada.pedidos:
            bruto = origem[p["id"]]
            for coluna, indice in mapa.items():
                campo = {"Nome do Cliente": "cliente", "Regiao": "regiao",
                         "Representante": "representante", "Status": "status"}[coluna]
                if bruto[indice] != p[campo]:
                    self.assertIn((p["id"], coluna), registradas,
                                  f"pedido {p['id']}: mudanca em {coluna} fora do log")

    def test_toda_pendencia_marcada_na_aba_principal(self):
        """Quem tem pendencia esta marcado como 'Revisar', e so quem tem"""
        com_pendencia = {r["id"] for r in self.gerada.revisar}
        marcados = {p["id"] for p in self.gerada.pedidos if p["situacao"] == "Revisar"}
        self.assertEqual(com_pendencia, marcados)
        for p in self.gerada.pedidos:
            self.assertIn(p["situacao"], {"OK", "Revisar"})

    def test_execucao_deterministica(self):
        """Duas execucoes produzem exatamente o mesmo dado"""
        segundo = self.tmp / "segunda.xlsx"
        pedidos = cp.carregar(ORIGINAL, cp.ABA_ORIGEM)
        pedidos, aud = cp.processar(pedidos)
        cp.gravar(segundo, pedidos, aud, ORIGINAL, len(pedidos))
        outra = Planilha(segundo)

        self.assertEqual(self.gerada.pedidos, outra.pedidos)
        self.assertEqual(self.gerada.revisar, outra.revisar)
        self.assertEqual(self.gerada.log, outra.log)
        self.assertEqual(self.gerada.removidos, outra.removidos)


# ===========================================================================
# 4. ESTRUTURA DA PLANILHA ENTREGUE
# ===========================================================================

class TestEstruturaDaPlanilha(BaseComPlanilha):

    def test_abas_esperadas(self):
        """A planilha traz resumo, dados, revisar, log e removidos"""
        self.assertEqual(set(self.gerada.abas),
                         {"resumo", "pedidos_vendas", "revisar", "log_correcoes", "removidos"})

    def test_cabecalho(self):
        """O cabecalho tem as colunas originais mais Situacao e Pendencias"""
        self.assertEqual(self.gerada.cabecalho,
                         ["ID Pedido", "Nome do Cliente", "Regiao", "Representante",
                          "Data do Pedido", "Receita", "Produto", "Status",
                          "Observacoes", "Situacao", "Pendencias"])

    def test_data_formatada(self):
        """A coluna de data usa o formato brasileiro"""
        coluna = self.gerada.cabecalho.index("Data do Pedido") + 1
        self.assertEqual(self.gerada.formato(2, coluna), cp.FORMATO_DATA)
        for p in self.gerada.pedidos:
            self.assertIsInstance(p["data"], (date, datetime))

    def test_listas_suspensas(self):
        """Regiao, Representante, Produto e Status tem lista suspensa"""
        validacoes = self.gerada.ws.data_validations.dataValidation
        colunas = {str(dv.sqref)[0] for dv in validacoes}
        self.assertTrue({"C", "D", "G", "H"} <= colunas,
                        f"faltam listas suspensas; presentes: {colunas}")

        status = next(dv for dv in validacoes if str(dv.sqref).startswith("H"))
        for valor in STATUS:
            self.assertIn(valor, status.formula1)

    def test_painel_congelado_e_filtro(self):
        """O cabecalho fica congelado e a aba tem filtro automatico"""
        self.assertEqual(self.gerada.ws.freeze_panes, "A2")
        self.assertIsNotNone(self.gerada.ws.auto_filter.ref)

    def test_resumo_tem_indicadores(self):
        """A aba resumo traz os indicadores de fechamento"""
        ws = self.gerada.livro["resumo"]
        rotulos = {linha[0] for linha in ws.iter_rows(min_row=2, values_only=True) if linha[0]}
        for esperado in ("Linhas lidas da origem", "Correcoes automaticas aplicadas",
                         "Linhas que ainda exigem revisao", "Receita ganha (apos correcoes)"):
            self.assertTrue(any(esperado in str(r) for r in rotulos), f"falta '{esperado}'")


# ===========================================================================
# 5. INTERFACE DE LINHA DE COMANDO
# ===========================================================================

class TestLinhaDeComando(unittest.TestCase):

    def setUp(self) -> None:
        if not ORIGINAL.exists():
            self.skipTest("planilha original ausente")
        self._dir = tempfile.mkdtemp(prefix="teste_cli_")
        self.tmp = Path(self._dir)

    def tearDown(self) -> None:
        shutil.rmtree(self._dir, ignore_errors=True)

    def executar(self, *args: str) -> subprocess.CompletedProcess:
        return subprocess.run([sys.executable, str(SCRIPT), *args],
                              capture_output=True, text=True, cwd=RAIZ)

    def test_execucao_normal(self):
        """A CLI roda e grava o arquivo pedido"""
        saida = self.tmp / "saida.xlsx"
        r = self.executar("--entrada", str(ORIGINAL), "--saida", str(saida))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue(saida.exists())

    def test_recusa_sobrescrever(self):
        """Sem --sobrescrever, a CLI se recusa a apagar um arquivo existente"""
        saida = self.tmp / "ja_existe.xlsx"
        saida.write_bytes(b"conteudo anterior")
        r = self.executar("--entrada", str(ORIGINAL), "--saida", str(saida))
        self.assertNotEqual(r.returncode, 0)
        self.assertEqual(saida.read_bytes(), b"conteudo anterior")

    def test_protege_o_arquivo_de_origem(self):
        """Mesmo com --sobrescrever, a CLI recusa gravar sobre a entrada"""
        antes = sha256(ORIGINAL)
        r = self.executar("--entrada", str(ORIGINAL), "--saida", str(ORIGINAL), "--sobrescrever")
        self.assertNotEqual(r.returncode, 0)
        self.assertEqual(sha256(ORIGINAL), antes)

    def test_entrada_inexistente(self):
        """Arquivo de entrada ausente falha com mensagem, nao com stack trace"""
        r = self.executar("--entrada", str(self.tmp / "nada.xlsx"),
                          "--saida", str(self.tmp / "s.xlsx"))
        self.assertNotEqual(r.returncode, 0)
        self.assertNotIn("Traceback", r.stdout + r.stderr)

    def test_aba_inexistente(self):
        """Aba inexistente falha com mensagem, nao com stack trace"""
        r = self.executar("--entrada", str(ORIGINAL), "--saida", str(self.tmp / "s.xlsx"),
                          "--aba", "inexistente")
        self.assertNotEqual(r.returncode, 0)
        self.assertNotIn("Traceback", r.stdout + r.stderr)


if __name__ == "__main__":
    print(f"\n  Script sob teste: {SCRIPT.name}")
    print(f"  Planilha de origem: {ORIGINAL}\n")
    unittest.main(verbosity=2)
