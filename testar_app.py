#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Suite de testes da camada web (app.py e planilha_io.py).

Fica separada de testar_correcoes.py de proposito: aquela cobre o motor de
limpeza e roda sem nenhuma dependencia alem de openpyxl; esta exige streamlit e
pandas, e cobre leitura de arquivo enviado, deteccao de colunas, exportacao e a
renderizacao da interface.

A interface e exercitada com o AppTest do streamlit. Como o AppTest nao sabe
simular upload, os testes executam o app por um arranjo que troca
st.file_uploader por um arquivo fixo - assim toda a tela de resultado (metricas,
abas, graficos, botoes de download) e de fato renderizada.

Uso
---
    python3 testar_app.py
    python3 testar_app.py -q
    python3 testar_app.py TestEquivalenciaCsvXlsx
"""

from __future__ import annotations

import csv
import io
from collections import Counter
import os
import shutil
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ))

import corrigir_planilha as cp          # noqa: E402
import limpeza_pandas as lp             # noqa: E402
import planilha_io as pio               # noqa: E402
import validacoes as vl                 # noqa: E402

APP = RAIZ / "app.py"
ORIGINAL = RAIZ / "data" / "planilha_vendas.xlsx"

CABECALHO_ALTERNATIVO = ["Nº Pedido", "Cliente", "Região", "Vendedor", "Data da Venda",
                         "Valor Total (R$)", "Item", "Situação", "Obs"]

ARRANJO = '''
import os, sys, streamlit as st
sys.path.insert(0, os.path.dirname(os.environ["APP"]))

class UploadFalso:
    def __init__(self, caminho):
        self.name = os.path.basename(caminho)
        self._dados = open(caminho, "rb").read()
    def getvalue(self):
        return self._dados

_arquivo = UploadFalso(os.environ["ARQUIVO_TESTE"])
st.file_uploader = lambda *a, **k: _arquivo

_app = os.environ["APP"]

# __file__ aponta para o app, nao para este arranjo: e dele que app.py resolve o
# proprio diretorio para achar os arquivos vizinhos.
_globais = {"__name__": "__main__", "__file__": _app}
exec(compile(open(_app, encoding="utf-8").read(), _app, "exec"), _globais)
'''


def linhas_da_origem() -> list:
    ws = openpyxl.load_workbook(ORIGINAL)["pedidos_vendas"]
    return list(ws.iter_rows(min_row=2, values_only=True))


def montar_csv(codec: str = "cp1252", separador: str = ";", decimal: str = ",",
               cabecalho: list = None) -> bytes:
    """Reproduz a planilha original como CSV, no estilo do Excel brasileiro."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=separador, lineterminator="\r\n")
    escritor.writerow(cabecalho or CABECALHO_ALTERNATIVO)
    for l in linhas_da_origem():
        receita = ""
        if l[5]:
            receita = str(l[5]).replace("R$ ", "").replace(".", decimal)
        escritor.writerow([
            "" if l[0] is None else int(l[0]), l[1] or "", l[2] or "", l[3] or "",
            l[4].strftime("%d/%m/%Y") if isinstance(l[4], datetime) else (l[4] or ""),
            receita, l[6] or "", l[7] or "", l[8] or "",
        ])
    return buffer.getvalue().encode(codec, errors="replace")


def limpar(registros: list, mapa: dict):
    pedidos = cp.pedidos_de_registros(registros, mapa)
    retrato = vl.Retrato.de(pedidos)
    limpos, aud = cp.processar(pedidos)
    return retrato, limpos, aud


# ===========================================================================
# LEITURA DO ARQUIVO ENVIADO
# ===========================================================================

class TestLeituraDeArquivo(unittest.TestCase):

    def test_detecta_cp1252(self):
        """A codificacao do Excel brasileiro e reconhecida, com acento intacto"""
        dados = montar_csv(codec="cp1252")
        registros, _, codec, _ = pio.ler_csv(dados)
        self.assertIn(codec, {"cp1252", "latin-1"})
        self.assertTrue(any("ã" in str(r) or "í" in str(r) for r in registros),
                        "a acentuacao se perdeu na leitura")

    def test_detecta_utf8_com_bom(self):
        """CSV em utf-8 com BOM e lido sem o marcador vazar para o cabecalho"""
        dados = montar_csv(codec="utf-8-sig")
        registros, cabecalho, codec, _ = pio.ler_csv(dados)
        self.assertEqual(codec, "utf-8-sig")
        self.assertFalse(cabecalho[0].startswith("﻿"))
        self.assertEqual(len(registros), len(linhas_da_origem()))

    def test_detecta_separadores(self):
        """Ponto e virgula, virgula e tabulacao sao todos reconhecidos"""
        for separador in (";", ",", "\t"):
            with self.subTest(separador=repr(separador)):
                dados = montar_csv(codec="utf-8", separador=separador, decimal=".")
                registros, cabecalho, _, achado = pio.ler_csv(dados)
                self.assertEqual(achado, separador)
                self.assertEqual(len(cabecalho), len(CABECALHO_ALTERNATIVO))
                self.assertEqual(len(registros), len(linhas_da_origem()))

    def test_separador_manual_vence_a_deteccao(self):
        """Quando o usuario informa codificacao e separador, eles sao respeitados"""
        dados = montar_csv(codec="utf-8", separador=";")
        _, _, codec, separador = pio.ler_csv(dados, codec="utf-8", separador=";")
        self.assertEqual((codec, separador), ("utf-8", ";"))

    def test_le_xlsx(self):
        """A planilha Excel e lida pelo cabecalho da primeira linha"""
        registros, cabecalho, abas = pio.ler_xlsx(ORIGINAL.read_bytes())
        self.assertEqual(abas, ["pedidos_vendas"])
        self.assertEqual(cabecalho[0], "ID Pedido")
        self.assertEqual(len(registros), len(linhas_da_origem()))

    def test_csv_so_com_cabecalho(self):
        """Arquivo sem linha de dados devolve lista vazia, sem estourar"""
        dados = "a;b;c\r\n".encode("utf-8")
        registros, cabecalho, _, _ = pio.ler_csv(dados)
        self.assertEqual(registros, [])
        self.assertEqual(cabecalho, ["a", "b", "c"])


# ===========================================================================
# DETECCAO DE COLUNAS
# ===========================================================================

class TestDeteccaoDeColunas(unittest.TestCase):

    def test_titulos_canonicos(self):
        """Os titulos da planilha original sao todos reconhecidos"""
        _, cabecalho, _ = pio.ler_xlsx(ORIGINAL.read_bytes())
        mapa = cp.detectar_colunas(cabecalho)
        self.assertEqual([c for c in pio.CAMPOS if not mapa.get(c)], [])

    def test_titulos_alternativos(self):
        """Titulos diferentes ('Vendedor', 'Valor Total (R$)') tambem sao reconhecidos"""
        mapa = cp.detectar_colunas(CABECALHO_ALTERNATIVO)
        self.assertEqual(mapa["id"], "Nº Pedido")
        self.assertEqual(mapa["representante"], "Vendedor")
        self.assertEqual(mapa["receita"], "Valor Total (R$)")
        self.assertEqual(mapa["produto"], "Item")
        self.assertEqual(mapa["status"], "Situação")
        self.assertEqual(mapa["observacoes"], "Obs")

    def test_nenhum_titulo_usado_duas_vezes(self):
        """Uma mesma coluna do arquivo nao e associada a dois campos"""
        mapa = cp.detectar_colunas(CABECALHO_ALTERNATIVO)
        usados = [v for v in mapa.values() if v]
        self.assertEqual(len(usados), len(set(usados)))

    def test_coluna_ausente_fica_nula(self):
        """Coluna que nao existe no arquivo e reportada como nao encontrada"""
        mapa = cp.detectar_colunas(["ID Pedido", "Cliente", "Data", "Valor"])
        self.assertIsNone(mapa["observacoes"])
        self.assertIsNone(mapa["status"])
        self.assertTrue(all(mapa[c] for c in pio.OBRIGATORIAS))

    def test_cabecalho_irreconhecivel(self):
        """Titulos sem relacao nenhuma nao produzem associacao falsa"""
        mapa = cp.detectar_colunas(["aaa", "bbb", "ccc"])
        self.assertEqual([v for v in mapa.values() if v], [])


# ===========================================================================
# EQUIVALENCIA ENTRE CSV E XLSX
# ===========================================================================

class TestEquivalenciaCsvXlsx(unittest.TestCase):
    """O mesmo dado, entregue como CSV ou como Excel, tem de dar no mesmo.

    Foi aqui que apareceu um defeito real: via CSV o ID chega como texto
    ('1001'), nao como decimal, e a checagem de cronologia - que filtrava por
    inteiro - era pulada em silencio.
    """

    @classmethod
    def setUpClass(cls):
        if not ORIGINAL.exists():
            raise unittest.SkipTest("planilha original ausente")
        registros_x, cab_x, _ = pio.ler_xlsx(ORIGINAL.read_bytes())
        cls.rx, cls.lx, cls.ax = limpar(registros_x, cp.detectar_colunas(cab_x))

        registros_c, cab_c, _, _ = pio.ler_csv(montar_csv())
        cls.rc, cls.lc, cls.ac = limpar(registros_c, cp.detectar_colunas(cab_c))

    def test_mesmas_linhas(self):
        """A contagem de linhas mantidas e removidas coincide"""
        self.assertEqual(len(self.lx), len(self.lc))
        self.assertEqual(len(self.ax.removidos), len(self.ac.removidos))

    def test_mesmos_campos_limpos(self):
        """Campo a campo, o resultado da limpeza e identico"""
        for px, pc in zip(sorted(self.lx, key=lambda p: p.id),
                          sorted(self.lc, key=lambda p: p.id)):
            with self.subTest(pedido=px.id):
                for campo in ("id", "cliente", "regiao", "representante",
                              "data", "receita", "produto", "status", "observacoes"):
                    self.assertEqual(getattr(px, campo), getattr(pc, campo), campo)

    def test_mesmas_correcoes_e_pendencias(self):
        """Pendencias identicas, e correcoes identicas fora a de tipo do ID

        A contagem de correcoes em 'ID Pedido' depende de como a origem entrega o
        campo: lido pelo pandas, o Excel ja devolve inteiro e nao ha o que
        corrigir, enquanto todo CSV entrega texto. O que precisa ser igual e o
        dado limpo, e disso cuida test_mesmos_campos_limpos.
        """
        def fora_do_id(aud):
            return Counter(c.coluna for c in aud.correcoes if c.coluna != "ID Pedido")

        self.assertEqual(fora_do_id(self.ax), fora_do_id(self.ac))
        self.assertEqual(len(self.ax.pendencias), len(self.ac.pendencias))
        self.assertEqual({p.tipo for p in self.ax.pendencias},
                         {p.tipo for p in self.ac.pendencias})

    def test_cronologia_verificada_nos_dois_caminhos(self):
        """A checagem de cronologia roda tambem quando o ID vem como texto"""
        for rotulo, aud in (("xlsx", self.ax), ("csv", self.ac)):
            with self.subTest(origem=rotulo):
                self.assertTrue(
                    any(p.tipo == "ID fora de ordem cronologica" for p in aud.pendencias),
                    "a cronologia nao foi verificada nesta origem")

    def test_id_sempre_inteiro(self):
        """O ID vira inteiro venha ele como decimal (xlsx) ou texto (csv)"""
        for rotulo, limpos in (("xlsx", self.lx), ("csv", self.lc)):
            with self.subTest(origem=rotulo):
                self.assertEqual({type(p.id).__name__ for p in limpos}, {"int"})

    def test_validacao_aprova_os_dois(self):
        """As 16 verificacoes passam nas duas origens"""
        for rotulo, (r, l, a) in (("xlsx", (self.rx, self.lx, self.ax)),
                                  ("csv", (self.rc, self.lc, self.ac))):
            with self.subTest(origem=rotulo):
                resumo = vl.resumir(vl.validar(r, l, a))
                self.assertTrue(resumo["tudo_ok"],
                                f"{resumo['erros']} erro(s), {resumo['alertas']} alerta(s)")


# ===========================================================================
# ID TEXTUAL
# ===========================================================================

class TestNormalizacaoDeId(unittest.TestCase):

    def test_formas_numericas_convergem(self):
        """1001, 1001.0 e '1001' viram o mesmo inteiro"""
        for entrada in (1001, 1001.0, "1001", " 1001 ", "1001.0"):
            with self.subTest(entrada=repr(entrada)):
                self.assertEqual(cp.ler_id(entrada), 1001)
                self.assertIsInstance(cp.ler_id(entrada), int)

    def test_id_textual_preservado(self):
        """Um ID genuinamente textual nao e destruido"""
        self.assertEqual(cp.ler_id("PED-001"), "PED-001")
        self.assertEqual(cp.ler_id("2024/A/17"), "2024/A/17")

    def test_vazio(self):
        """Vazio e nulo continuam nulos"""
        self.assertIsNone(cp.ler_id(None))
        self.assertIsNone(cp.ler_id("   "))


# ===========================================================================
# CAMADA PANDAS
# ===========================================================================

class TestCamadaPandas(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        if not ORIGINAL.exists():
            raise unittest.SkipTest("planilha original ausente")
        cls.bruto = lp.ler(ORIGINAL)
        cls.r = lp.limpar(cls.bruto)

    def test_ler_devolve_dataframe_bruto(self):
        """ler() entrega o arquivo como veio, sem nenhuma correcao"""
        self.assertIsInstance(self.bruto, pd.DataFrame)
        self.assertEqual(len(self.bruto), len(linhas_da_origem()))
        self.assertIn("R$ 4320.00", self.bruto["Receita"].tolist(),
                      "a receita ja chegou alterada; ler() nao pode limpar")

    def test_ler_aceita_caminho_e_bytes(self):
        """ler() aceita caminho, bytes e objeto com getvalue()"""
        por_caminho = lp.ler(ORIGINAL)
        por_bytes = lp.ler(montar_csv())
        self.assertEqual(len(por_caminho), len(por_bytes))

    def test_tipos_do_dataframe_limpo(self):
        """O DataFrame limpo sai com data como datetime e receita como float"""
        self.assertEqual(str(self.r.df["Data do Pedido"].dtype), "datetime64[ns]")
        self.assertEqual(str(self.r.df["Receita"].dtype), "float64")
        self.assertTrue(pd.api.types.is_integer_dtype(self.r.df["ID Pedido"]))
        self.assertEqual(list(self.r.df.columns), pio.COLUNAS_SAIDA)

    def test_resultado_expoe_tudo_em_dataframe(self):
        """correcoes, pendencias, removidos e laudo sao DataFrames"""
        for nome in ("correcoes", "pendencias", "removidos", "laudo"):
            with self.subTest(atributo=nome):
                self.assertIsInstance(getattr(self.r, nome), pd.DataFrame)
        self.assertEqual(len(self.r.correcoes), len(self.r.auditoria.correcoes))
        self.assertEqual(len(self.r.pendencias), len(self.r.auditoria.pendencias))
        self.assertEqual(self.r.total_problemas,
                         self.r.total_melhorias + self.r.total_pendencias)

    def test_equivale_ao_motor_linha_a_linha(self):
        """A camada pandas devolve o mesmo que o motor de regras direto"""
        registros, cabecalho, _ = pio.ler_xlsx(ORIGINAL.read_bytes())
        _, limpos, aud = limpar(registros, cp.detectar_colunas(cabecalho))
        self.assertEqual(len(self.r.df), len(limpos))
        self.assertEqual(self.r.df["ID Pedido"].tolist(), [p.id for p in limpos])
        self.assertAlmostEqual(float(self.r.df["Receita"].sum()),
                               sum(p.receita or 0 for p in limpos), places=2)
        self.assertEqual(len(self.r.correcoes), len(aud.correcoes))

    def test_metricas(self):
        """As metricas reproduzem a receita ganha apurada no diagnostico"""
        m = lp.metricas(self.r.df)
        self.assertAlmostEqual(m["receita_ganha"], 113_720.00, places=2)
        self.assertEqual(m["sem_receita"], 2)
        self.assertEqual(m["ganhos"] + m["perdidos"], len(self.r.df))
        # Denominador e o de ganhos COM receita: os sem valor nao entram como zero.
        self.assertEqual(m["ganhos_com_receita"], m["ganhos"] - m["sem_receita"])
        self.assertAlmostEqual(m["ticket_medio"],
                               m["receita_ganha"] / m["ganhos_com_receita"], places=2)
        self.assertTrue(0 <= m["taxa_de_ganho"] <= 1)

    def test_metricas_de_dataframe_vazio(self):
        """Arquivo sem linha nao quebra o calculo de metricas"""
        m = lp.metricas(pd.DataFrame(columns=pio.COLUNAS_SAIDA))
        self.assertEqual(m["pedidos"], 0)
        self.assertEqual(m["receita_ganha"], 0.0)

    def test_resumo_por_dimensao(self):
        """O agrupamento fecha com o total e cobre todas as categorias"""
        for dimensao in ("Regiao", "Representante", "Produto"):
            with self.subTest(dimensao=dimensao):
                tabela = lp.resumo_por(self.r.df, dimensao)
                self.assertEqual(len(tabela), self.r.df[dimensao].nunique())
                self.assertEqual(int(tabela["Pedidos"].sum()), len(self.r.df))
                self.assertAlmostEqual(float(tabela["Receita ganha"].sum()),
                                       lp.metricas(self.r.df)["receita_ganha"], places=2)
                self.assertTrue((tabela["Taxa de ganho"] <= 1).all())

    def test_serie_mensal(self):
        """A serie mensal soma o mesmo total e vem em ordem cronologica"""
        mensal = lp.serie_mensal(self.r.df)
        self.assertAlmostEqual(float(mensal["Receita ganha"].sum()),
                               lp.metricas(self.r.df)["receita_ganha"], places=2)
        self.assertEqual(mensal["Mes"].tolist(), sorted(mensal["Mes"].tolist()))
        self.assertRegex(mensal["Mes"].iloc[0], r"^\d{4}-\d{2}$")

    def test_qualidade_e_pendencias_agregadas(self):
        """As agregacoes de qualidade fecham com os totais brutos"""
        qualidade = lp.qualidade_por_coluna(self.r)
        self.assertEqual(int(qualidade["Correcoes"].sum()), self.r.total_melhorias)
        tipos = lp.pendencias_por_tipo(self.r)
        self.assertEqual(int(tipos["Ocorrencias"].sum()), self.r.total_pendencias)

    def test_limpar_arquivo_e_atalho(self):
        """limpar_arquivo() equivale a ler() seguido de limpar()"""
        direto = lp.limpar_arquivo(ORIGINAL)
        self.assertEqual(len(direto.df), len(self.r.df))
        self.assertEqual(direto.total_melhorias, self.r.total_melhorias)


# ===========================================================================
# EXPORTACAO
# ===========================================================================

class TestExportacao(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        if not ORIGINAL.exists():
            raise unittest.SkipTest("planilha original ausente")
        registros, cabecalho, _ = pio.ler_xlsx(ORIGINAL.read_bytes())
        cls.retrato, cls.limpos, cls.aud = limpar(registros, cp.detectar_colunas(cabecalho))

    def test_csv_perfil_brasileiro(self):
        """O CSV sai com ponto e virgula, virgula decimal e BOM para o Excel"""
        sep, dec, codec = pio.PERFIS_CSV["Excel brasileiro"]
        dados = pio.para_csv(self.limpos, sep, dec, codec)
        self.assertTrue(dados.startswith(b"\xef\xbb\xbf"), "falta o BOM do utf-8-sig")
        texto = dados.decode("utf-8-sig")
        self.assertEqual(texto.splitlines()[0].split(";"), pio.COLUNAS_SAIDA)
        self.assertIn("4320,00", texto)
        self.assertNotIn("4320.00", texto)

    def test_csv_perfil_internacional(self):
        """O perfil internacional usa virgula e ponto decimal"""
        sep, dec, codec = pio.PERFIS_CSV["Padrao internacional"]
        texto = pio.para_csv(self.limpos, sep, dec, codec).decode("utf-8")
        self.assertEqual(texto.splitlines()[0].split(","), pio.COLUNAS_SAIDA)
        self.assertIn("4320.00", texto)

    def test_csv_relido_bate_com_a_origem(self):
        """O CSV exportado volta a ser lido com o mesmo numero de linhas e valores"""
        sep, dec, codec = pio.PERFIS_CSV["Excel brasileiro"]
        dados = pio.para_csv(self.limpos, sep, dec, codec)
        registros, _, _, _ = pio.ler_csv(dados)
        self.assertEqual(len(registros), len(self.limpos))
        total = sum(cp.ler_receita(r["Receita"]) or 0 for r in registros)
        esperado = sum(p.receita or 0 for p in self.limpos)
        self.assertAlmostEqual(total, esperado, places=2)

    def test_csv_sem_none_literal(self):
        """Celula vazia sai em branco, nunca como a palavra None"""
        texto = pio.para_csv(self.limpos).decode("utf-8-sig")
        self.assertNotIn("None", texto)

    def test_excel_tem_as_cinco_abas(self):
        """O Excel exportado traz as cinco abas, com o mesmo numero de linhas"""
        dados = pio.para_excel(self.limpos, self.aud, "origem.xlsx", self.retrato.linhas)
        livro = openpyxl.load_workbook(io.BytesIO(dados))
        self.assertEqual(set(livro.sheetnames),
                         {"resumo", "pedidos_vendas", "revisar", "log_correcoes", "removidos"})
        self.assertEqual(livro["pedidos_vendas"].max_row - 1, len(self.limpos))
        self.assertEqual(livro["log_correcoes"].max_row - 1, len(self.aud.correcoes))
        self.assertEqual(livro["revisar"].max_row - 1, len(self.aud.pendencias))

    def test_data_em_formato_brasileiro(self):
        """A data sai como dd/mm/aaaa nas linhas exportadas"""
        for linha in pio.como_linhas(self.limpos):
            if linha["Data do Pedido"]:
                self.assertRegex(str(linha["Data do Pedido"]), r"^\d{2}/\d{2}/\d{4}$")


# ===========================================================================
# INTERFACE
# ===========================================================================

class TestInterface(unittest.TestCase):
    """Roda o app de verdade com o AppTest do streamlit."""

    @classmethod
    def setUpClass(cls):
        try:
            from streamlit.testing.v1 import AppTest  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("streamlit nao instalado")
        if not APP.exists():
            raise unittest.SkipTest("app.py ausente")
        cls._dir = tempfile.mkdtemp(prefix="teste_ui_")
        cls.arranjo = Path(cls._dir) / "arranjo.py"
        cls.arranjo.write_text(ARRANJO, encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(getattr(cls, "_dir", ""), ignore_errors=True)

    def setUp(self):
        import streamlit
        self._uploader_original = streamlit.file_uploader

    def tearDown(self):
        # O arranjo troca st.file_uploader no processo em que o AppTest roda;
        # sem restaurar, o upload falso vaza para o teste seguinte.
        import streamlit
        streamlit.file_uploader = self._uploader_original

    def executar(self, arquivo: Path):
        """Roda o app com o arquivo ja enviado.

        A limpeza acontece assim que o arquivo chega: nao ha botao para clicar.
        """
        from streamlit.testing.v1 import AppTest
        os.environ["APP"] = str(APP)
        os.environ["ARQUIVO_TESTE"] = str(arquivo)
        at = AppTest.from_file(str(self.arranjo), default_timeout=300).run()
        self.assertFalse(at.exception, f"excecao: {[e.value for e in at.exception]}")
        return at

    def test_pagina_abre_sem_arquivo(self):
        """Sem upload, a pagina abre com a orientacao e a tabela de colunas"""
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_file(str(APP), default_timeout=120).run()
        self.assertFalse(at.exception)
        self.assertIn("Limpeza Automatizada", at.title[0].value)
        self.assertTrue(at.get("file_uploader"))
        self.assertTrue(at.info)
        self.assertEqual([t.label for t in at.sidebar.toggle],
                         ["Rodar suite de testes ao final",
                          "Imputar receita ausente pela mediana do produto"])

    def test_fluxo_completo_com_csv(self):
        """Com um CSV enviado, a tela de resultado e renderizada inteira"""
        csv_tmp = Path(self._dir) / "pedidos.csv"
        csv_tmp.write_bytes(montar_csv())
        at = self.executar(csv_tmp)

        topo = {m.label: m.value for m in at.metric[:4]}
        self.assertEqual(list(topo), ["Registros (entrada)", "Registros (saida)",
                                      "Erros corrigidos", "Pendencias para decidir"])
        self.assertEqual(topo["Registros (entrada)"], str(len(linhas_da_origem())))
        cabecalhos = [h.value for h in at.subheader]
        for esperado in ("Resumo", "Antes e depois", "Validacao automatizada"):
            self.assertTrue(any(esperado in c for c in cabecalhos),
                            f"secao '{esperado}' ausente: {cabecalhos}")

        recolhidos = [e.label for e in at.expander]
        for esperado in ("Correcoes aplicadas", "Pendencias para voce decidir",
                         "Detalhe das verificacoes", "Dados limpos",
                         "Analise da planilha limpa"):
            self.assertTrue(any(esperado in c for c in recolhidos),
                            f"secao recolhida '{esperado}' ausente: {recolhidos}")

        self.assertEqual([b.label for b in at.get("download_button")],
                         ["📄 Baixar CSV", "📊 Baixar Excel"])

    def test_fluxo_completo_com_xlsx(self):
        """O mesmo fluxo funciona com upload de Excel"""
        at = self.executar(ORIGINAL)
        self.assertTrue(
            any("Limpeza concluida e aprovada" in s.value for s in at.success),
            f"a tela nao confirmou a limpeza: {[s.value for s in at.success]}")

    def test_metricas_batem_com_o_motor(self):
        """Os numeros na tela sao os mesmos que o motor de limpeza produz"""
        registros, cabecalho, _ = pio.ler_xlsx(ORIGINAL.read_bytes())
        _, limpos, aud = limpar(registros, cp.detectar_colunas(cabecalho))

        at = self.executar(ORIGINAL)
        topo = {m.label: m.value for m in at.metric[:4]}
        self.assertEqual(topo["Registros (entrada)"], str(len(linhas_da_origem())))
        self.assertEqual(topo["Registros (saida)"], str(len(limpos)))
        self.assertEqual(topo["Erros corrigidos"], str(len(aud.correcoes)))
        self.assertEqual(topo["Pendencias para decidir"], str(len(aud.pendencias)))
        self.assertTrue(
            any(f"({len(aud.correcoes)})" in e.label for e in at.expander
                if "Correcoes aplicadas" in e.label),
            f"o rotulo recolhido nao traz a contagem: {[e.label for e in at.expander]}")

    def test_suite_de_regressao_roda_ao_final(self):
        """A suite roda e reporta contagem e tempo, no formato 'N passed in Xs'"""
        at = self.executar(ORIGINAL)
        mensagens = [s.value for s in at.success]
        achou = next((m for m in mensagens if "Todos os testes passaram" in m), None)
        self.assertIsNotNone(achou, f"a suite nao reportou sucesso: {mensagens}")
        self.assertRegex(achou, r"\d+ passed in \d+\.\d{2}s",
                         "o resultado da suite nao traz contagem e tempo")

    def test_validacao_fica_visivel_fora_de_expansor(self):
        """O resultado da validacao aparece na pagina, nao escondido num expansor"""
        at = self.executar(ORIGINAL)
        self.assertTrue(any("Validacao automatizada" in h.value for h in at.subheader),
                        [h.value for h in at.subheader])

    def test_antes_e_depois_mostram_a_melhora(self):
        """O comparativo mostra os defeitos caindo entre antes e depois"""
        bruto = lp.ler(ORIGINAL)
        r = lp.limpar(bruto)
        antes = lp.diagnosticar(bruto, r.mapa)
        depois = lp.diagnosticar(r.df)
        for chave in ("receita_texto", "status_fora", "caixa_inconsistente", "nulos"):
            with self.subTest(indicador=chave):
                self.assertLess(depois[chave], antes[chave],
                                f"{chave} nao melhorou: {antes[chave]} -> {depois[chave]}")
        tabela = lp.comparativo(antes, depois)
        self.assertEqual(list(tabela.columns),
                         ["Indicador", "Antes", "Depois", "Variacao", "Situacao"])
        self.assertTrue((tabela["Variacao"] <= 0).all(),
                        "algum indicador de defeito piorou com a limpeza")
        self.assertNotIn("Linhas", tabela["Indicador"].tolist(),
                         "contagem de linhas nao e defeito e nao deve entrar aqui")

        at = self.executar(ORIGINAL)
        self.assertTrue(any("Antes e depois" in h.value for h in at.subheader),
                        [h.value for h in at.subheader])


if __name__ == "__main__":
    print(f"\n  App sob teste: {APP.name}")
    print(f"  Planilha de referencia: {ORIGINAL}\n")
    unittest.main(verbosity=2)
